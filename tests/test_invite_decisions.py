from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from test_invite_portal import ROOT, finish, prepare, row
from test_invite_portal import admin as admin_fixture
from test_invite_portal import portal as portal_fixture

from app.extensions import db
from app.models import InvitationPortalGuest

admin = admin_fixture
portal = portal_fixture


@pytest.mark.parametrize("prefix", ["/admin", "/u/groom-token"])
def test_question_blocks_all_send_paths(admin, portal, prefix):
    attempt = prepare(admin, prefix=prefix).json["attempt"]
    url = f"{ROOT}{prefix}/guest/1"
    assert (
        admin.post(url + "/decision", data={"invitation_decision": "undecided"}).status_code == 302
    )
    assert row(portal)[0] == "unsent"
    assert prepare(admin, prefix=prefix).status_code == 409
    assert finish(admin, attempt, prefix=prefix).status_code == 409
    for status in ("preparing", "sent"):
        assert admin.post(url + "/status", data={"status": status}).status_code == 409
    assert admin.post(url + "/decision", data={"invitation_decision": "invalid"}).status_code == 400
    assert admin.post(url + "/decision", data={"invitation_decision": "invited"}).status_code == 302
    assert prepare(admin, prefix=prefix).status_code == 200


@pytest.mark.parametrize("prefix", ["/admin", "/u/groom-token"])
def test_delete_restore_pending_guest(admin, portal, prefix):
    attempt = prepare(admin, prefix=prefix).json["attempt"]
    url = f"{ROOT}{prefix}/guest/1"
    assert admin.post(url + "/delete").status_code == 302
    assert row(portal)[0] == "unsent"
    assert 'data-guest-row="1"' not in admin.get(ROOT + prefix).text
    assert 'data-guest-row="1"' in admin.get(ROOT + prefix + "/trash").text
    assert admin.get(url + "/edit").status_code == 404
    assert prepare(admin, prefix=prefix).status_code == 404
    assert finish(admin, attempt, prefix=prefix).status_code == 404
    for action, data in [
        ("decision", {"invitation_decision": "invited"}),
        ("status", {"status": "sent"}),
    ]:
        assert admin.post(url + "/" + action, data=data).status_code == 404
    book = load_workbook(BytesIO(admin.get(ROOT + prefix + "/export.xlsx").data))
    assert "משה" not in [r[1] for r in book.active.iter_rows(min_row=2, values_only=True)]
    assert admin.post(url + "/restore").status_code == 302
    assert 'data-guest-row="1"' in admin.get(ROOT + prefix).text
    assert finish(admin, attempt, prefix=prefix).status_code == 409


def test_restore_preserves_question_and_sent_history(admin, portal):
    attempt = prepare(admin).json["attempt"]
    assert finish(admin, attempt).status_code == 200
    before = row(portal)
    url = ROOT + "/admin/guest/1"
    admin.post(url + "/decision", data={"invitation_decision": "undecided"})
    admin.post(url + "/delete")
    admin.post(url + "/restore")
    assert row(portal) == before
    assert prepare(admin).status_code == 409
    filtered = admin.get(ROOT + "/admin?invitation_decision=undecided").text
    assert 'data-guest-row="1"' in filtered
    assert 'data-guest-row="2"' not in filtered


@pytest.mark.parametrize("action", ["decision", "delete", "restore"])
def test_mutations_scoped_and_authenticated(client, portal, action):
    assert client.post(f"{ROOT}/u/bride-token/guest/1/{action}").status_code == 404
    assert client.post(f"{ROOT}/admin/guest/1/{action}").status_code in (302, 401)
    client.post(f"{ROOT}/u/groom-token/guest/1/delete")
    assert 'data-guest-row="1"' not in client.get(ROOT + "/u/bride-token/trash").text


def test_excel_decision_roundtrip_and_deleted_rejection(admin, portal):
    admin.post(ROOT + "/admin/guest/1/decision", data={"invitation_decision": "undecided"})
    exported = admin.get(ROOT + "/admin/export.xlsx").data
    book = load_workbook(BytesIO(exported))
    assert book.active["M2"].value == "בסימן שאלה"
    assert (
        admin.post(
            ROOT + "/admin/import", data={"file": (BytesIO(exported), "guests.xlsx")}
        ).status_code
        == 302
    )
    assert prepare(admin).status_code == 409
    book.active.delete_cols(13)
    legacy = BytesIO()
    book.save(legacy)
    legacy.seek(0)
    assert (
        admin.post(ROOT + "/admin/import", data={"file": (legacy, "legacy.xlsx")}).status_code
        == 302
    )
    assert prepare(admin).status_code == 409
    admin.post(ROOT + "/admin/guest/1/delete")
    assert (
        admin.post(
            ROOT + "/admin/import", data={"file": (BytesIO(exported), "guests.xlsx")}
        ).status_code
        == 400
    )


def test_guest_edit_can_set_question(admin, portal):
    response = admin.post(
        ROOT + "/admin/guest/1/edit",
        data={
            "first_name": "משה",
            "side": "groom",
            "salutation": "male",
            "invitation_decision": "undecided",
        },
    )
    assert response.status_code == 302
    assert prepare(admin).status_code == 409


@pytest.mark.parametrize("status", ["sent", "preparing"])
def test_excel_cannot_mark_question_guest_as_sent(admin, portal, status):
    admin.post(ROOT + "/admin/guest/1/decision", data={"invitation_decision": "undecided"})
    with portal.app_context():
        identifier = db.session.get(InvitationPortalGuest, 1).external_id
    csv = f"מזהה,שם פרטי,צד,צורת פנייה,סטטוס\n{identifier},משה,groom,male,{status}\n"
    response = admin.post(
        ROOT + "/admin/import",
        data={
            "file": (BytesIO(csv.encode("utf-8-sig")), "guests.csv"),
        },
    )
    assert response.status_code == 400
    assert row(portal)[0] == "unsent"
    assert prepare(admin).status_code == 409


@pytest.mark.parametrize("action", ["decision", "delete", "restore"])
def test_new_mutations_require_csrf(admin, portal, action):
    portal.config["WTF_CSRF_ENABLED"] = True
    for prefix in ("/admin", "/u/groom-token"):
        assert admin.post(f"{ROOT}{prefix}/guest/1/{action}").status_code == 400


def test_init_db_upgrades_existing_guests_idempotently(portal):
    with portal.app_context():
        db.session.remove()
        with db.engine.begin() as connection:
            connection.execute(
                text("ALTER TABLE invitation_portal_guests DROP COLUMN invitation_decision")
            )
            connection.execute(text("ALTER TABLE invitation_portal_guests DROP COLUMN deleted_at"))
    for _ in range(2):
        result = portal.test_cli_runner().invoke(args=["init-db"])
        assert result.exit_code == 0, result.output
    with portal.app_context():
        guest = db.session.get(InvitationPortalGuest, 1)
        assert guest.first_name == "משה"
        assert guest.invitation_decision == "invited"
        assert guest.deleted_at is None

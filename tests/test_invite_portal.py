from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.extensions import db
from app.models import (
    Guest,
    InvitationPortalGuest,
    InvitationPortalSettings,
    InvitationSendAttempt,
    InvitationSender,
    User,
)

ROOT = "/invite-manager"
JSON = {"Accept": "application/json"}


@pytest.fixture()
def portal(app, tmp_path):
    app.config["UPLOAD_FOLDER"] = tmp_path
    app.config["RATELIMIT_ENABLED"] = False
    Image.new("RGB", (80, 120), "white").save(tmp_path / "test.jpg")
    with app.app_context():
        senders = [
            InvitationSender(
                name=name,
                role=role,
                side=side,
                access_token=token,
                male_template="{name} היקר",
                female_template="{name} היקרה",
                plural_template="{name} היקרים",
            )
            for name, role, side, token in [
                ("יוסי", "אבא של החתן", "groom", "groom-token"),
                ("מיכל", "אמא של הכלה", "bride", "bride-token"),
                ("דנה", "אמא של החתן", "groom", "other-token"),
            ]
        ]
        db.session.add_all(senders)
        db.session.add(InvitationPortalSettings(id=1, image_filename="test.jpg"))
        db.session.add_all(
            [
                InvitationPortalGuest(
                    first_name="משה",
                    last_name="לוי",
                    phone="0501234567",
                    side="groom",
                    salutation="male",
                    group_name="חברים",
                ),
                InvitationPortalGuest(
                    first_name="מרים",
                    last_name="כהן",
                    phone="0521234567",
                    side="bride",
                    salutation="female",
                    group_name="משפחה",
                ),
                InvitationPortalGuest(
                    first_name="יצחק והילה", phone="0541234567", side="groom", salutation="plural"
                ),
            ]
        )
        db.session.commit()
    return app


@pytest.fixture()
def admin(client, portal):
    response = client.post(
        f"{ROOT}/admin/login", data={"email": "admin@example.com", "password": "password123"}
    )
    assert response.status_code == 302
    return client


def prepare(client, guest_id=1, prefix="/u/groom-token", **data):
    return client.post(f"{ROOT}{prefix}/guest/{guest_id}/prepare", data=data, headers=JSON)


def finish(client, attempt, guest_id=1, prefix="/u/groom-token", sent="true"):
    return client.post(
        f"{ROOT}{prefix}/guest/{guest_id}/confirm",
        data={"attempt": attempt, "sent": sent},
        headers=JSON,
    )


def row(app, guest_id=1):
    with app.app_context():
        guest = db.session.get(InvitationPortalGuest, guest_id)
        return guest.status, guest.attempts, guest.last_sender_id, guest.sent_at


@pytest.mark.parametrize(
    "url",
    [
        "/admin",
        "/admin/senders",
        "/admin/settings",
        "/admin/guest/new",
        "/admin/guest/1/edit",
        "/admin/senders/1",
    ],
)
def test_admin_pages_render(admin, url):
    response = admin.get(ROOT + url)
    assert response.status_code == 200
    assert "רגע של שמחה".encode() in response.data
    assert b"app.css" not in response.data


def test_login_is_public_and_does_not_expose_admin_navigation(client, portal):
    response = client.get(ROOT + "/admin/login")
    assert response.status_code == 200
    assert b'href="/invite-manager/admin/senders"' not in response.data
    assert (
        client.post(
            ROOT + "/admin/login", data={"email": "admin@example.com", "password": "wrong"}
        ).status_code
        == 200
    )
    assert client.get(ROOT + "/admin").status_code == 302


@pytest.mark.parametrize(
    "path", ["/admin", "/admin/settings", "/admin/senders", "/admin/export.xlsx"]
)
def test_admin_endpoints_protected(client, portal, path):
    assert client.get(ROOT + path).status_code == 302


@pytest.mark.parametrize(
    "path",
    [
        "/admin/guest/1/prepare",
        "/admin/guest/1/confirm",
        "/admin/guest/1/status",
        "/admin/import",
        "/admin/senders/1/renew",
    ],
)
def test_admin_writes_protected(client, portal, path):
    assert client.post(ROOT + path, headers=JSON).status_code == 401


def test_sender_scope_and_revocation(admin, portal):
    response = admin.get(ROOT + "/u/groom-token")
    assert "משה".encode() in response.data
    assert "מרים".encode() not in response.data
    assert prepare(admin, 2).status_code == 404
    assert admin.get(ROOT + "/u/groom-token/guest/2/edit").status_code == 404
    assert (
        admin.post(ROOT + "/u/groom-token/guest/2/status", data={"status": "sent"}).status_code
        == 404
    )
    assert admin.post(ROOT + "/admin/senders/1/renew").status_code == 302
    assert admin.get(ROOT + "/u/groom-token").status_code == 404
    with portal.app_context():
        sender = db.session.get(InvitationSender, 1)
        token = sender.access_token
        sender.is_active = False
        db.session.commit()
    assert admin.get(ROOT + "/u/" + token).status_code == 404
    assert prepare(admin, prefix="/u/" + token).status_code == 404


@pytest.mark.parametrize(
    "guest_id,prefix,expected",
    [
        (1, "/u/groom-token", "משה היקר"),
        (2, "/u/bride-token", "מרים היקרה"),
        (3, "/u/groom-token", "יצחק והילה היקרים"),
        (1, "/admin", "משה היקר"),
    ],
)
def test_prepare_and_confirm_explicitly(admin, portal, guest_id, prefix, expected):
    response = prepare(admin, guest_id, prefix)
    assert response.status_code == 200
    data = response.json
    assert data["text"].startswith(expected)
    assert "לוי" not in data["text"]
    assert row(portal, guest_id)[:2] == ("preparing", 0)
    response = finish(admin, data["attempt"], guest_id, prefix)
    assert response.status_code == 200
    assert row(portal, guest_id)[0] == "sent"
    assert row(portal, guest_id)[1] == 1
    assert row(portal, guest_id)[3] is not None
    assert finish(admin, data["attempt"], guest_id, prefix).status_code == 200
    assert row(portal, guest_id)[1] == 1


def test_admin_can_select_sender_and_reject_wrong_side(admin, portal):
    assert prepare(admin, prefix="/admin", sender_id=2).status_code == 400
    response = prepare(admin, prefix="/admin", sender_id=1)
    assert response.json["text"] == "משה היקר"
    assert finish(admin, response.json["attempt"], prefix="/admin").status_code == 200
    assert row(portal)[2] == 1


def test_cancel_resend_restores_previous_status_and_history(client, portal):
    data = prepare(client).json
    finish(client, data["attempt"])
    before = row(portal)
    resend = prepare(client).json
    assert finish(client, resend["attempt"], sent="false").status_code == 200
    assert row(portal) == before
    assert finish(client, resend["attempt"]).status_code == 409


def test_cancel_first_send_and_resume_after_refresh(client, portal):
    data = prepare(client).json
    resumed = prepare(client).json
    assert resumed["resumed"]
    assert resumed["attempt"] == data["attempt"]
    assert finish(client, data["attempt"], sent="false").status_code == 200
    assert row(portal)[:2] == ("unsent", 0)


def test_two_senders_cannot_confirm_each_others_attempt(client, portal):
    data = prepare(client).json
    assert prepare(client, prefix="/u/other-token").status_code == 409
    assert finish(client, data["attempt"], prefix="/u/other-token").status_code == 403
    assert finish(client, data["attempt"], guest_id=3).status_code == 400


@pytest.mark.parametrize("prefix", ["/admin", "/u/groom-token"])
def test_send_with_only_first_name_and_salutation(admin, portal, prefix):
    with portal.app_context():
        guest = db.session.get(InvitationPortalGuest, 1)
        guest.phone = ""
        guest.last_name = ""
        db.session.commit()
    response = prepare(admin, prefix=prefix)
    assert response.status_code == 200
    assert response.json["whatsapp_phone"] is None
    assert response.json["text"].startswith("משה היקר")
    assert finish(admin, response.json["attempt"], prefix=prefix).status_code == 200
    assert row(portal)[0] == "sent"


def test_missing_image_is_explained(client, portal):
    with portal.app_context():
        db.session.get(InvitationPortalGuest, 1).phone = "0501234567"
        db.session.get(InvitationPortalSettings, 1).image_filename = "missing.jpg"
        db.session.commit()
    response = prepare(client)
    assert response.status_code == 400
    assert "תמונת" in response.json["error"]
    assert row(portal)[0] == "unsent"


@pytest.mark.parametrize("prefix", ["/admin", "/u/groom-token"])
def test_create_guest_without_optional_contact_fields(admin, portal, prefix):
    response = admin.post(
        ROOT + prefix + "/guest/new",
        data={"first_name": "דוד", "salutation": "male", "side": "groom"},
    )
    assert response.status_code == 302
    with portal.app_context():
        guest = db.session.scalar(
            db.select(InvitationPortalGuest).where(InvitationPortalGuest.first_name == "דוד")
        )
        assert guest is not None and not guest.last_name and not guest.phone
    assert "יצחק והילה".encode() not in admin.get(ROOT + prefix + "/guest/new").data


def test_confirmation_requires_preparation_and_valid_decision(client, portal):
    assert finish(client, "fake").status_code == 400
    data = prepare(client).json
    assert finish(client, data["attempt"], sent="maybe").status_code == 400
    assert row(portal)[0] == "preparing"


def test_manual_status_invalidates_stale_confirmation(admin, portal):
    data = prepare(admin).json
    admin.post(ROOT + "/admin/guest/1/status", data={"status": "unsent"})
    assert finish(admin, data["attempt"]).status_code == 409
    assert row(portal)[0] == "unsent"


def test_sender_can_edit_own_messages_and_guests(client, portal):
    response = client.post(
        ROOT + "/u/groom-token/message",
        data={
            "male_template": "{name} שלום",
            "female_template": "{name} שלום לך",
            "plural_template": "{name} שלום לכם",
        },
    )
    assert response.status_code == 302
    assert prepare(client).json["text"] == "משה שלום"
    response = client.post(
        ROOT + "/u/groom-token/guest/new",
        data={"first_name": "יעל", "side": "bride", "salutation": "female", "phone": "501234567"},
    )
    assert response.status_code == 302
    with portal.app_context():
        guest = db.session.scalar(
            db.select(InvitationPortalGuest).where(InvitationPortalGuest.first_name == "יעל")
        )
        assert guest.side == "groom"
        assert guest.phone == "0501234567"
        assert db.session.scalar(db.select(db.func.count(Guest.id))) == 0


@pytest.mark.parametrize("salutation", ["", "unknown"])
def test_salutation_is_mandatory(admin, portal, salutation):
    response = admin.post(
        ROOT + "/admin/guest/new",
        data={
            "first_name": "יצחק והילה",
            "side": "groom",
            "salutation": salutation,
        },
    )
    assert response.status_code == 200
    assert "יש לתקן".encode() in response.data
    with portal.app_context():
        assert db.session.scalar(db.select(db.func.count(InvitationPortalGuest.id))) == 3


def test_sender_crud_and_existing_list_render(admin):
    response = admin.post(
        ROOT + "/admin/senders",
        data={
            "name": "החתן",
            "role": "חתן",
            "side": "groom",
            "is_active": "y",
            "male_template": "{name} א",
            "female_template": "{name} ב",
            "plural_template": "{name} ג",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "החתן".encode() in response.data
    assert b"data-copy-target" in response.data


def test_admin_edits_each_senders_templates_without_changing_other_senders(admin, portal):
    for sender_id in [1, 2]:
        response = admin.get(ROOT + f"/admin/senders/{sender_id}")
        assert response.status_code == 200
        assert b'class="sender-list"' not in response.data
        assert "שמירת הנוסחים והפרטים".encode() in response.data
        with portal.app_context():
            sender = db.session.get(InvitationSender, sender_id)
            data = {
                "name": sender.name,
                "role": sender.role,
                "side": sender.side,
                "is_active": "y",
                "male_template": f"{{name}} זכר {sender_id}",
                "female_template": f"{{name}} נקבה {sender_id}",
                "plural_template": f"{{name}} רבים {sender_id}",
            }
            token = sender.access_token
        saved = admin.post(ROOT + f"/admin/senders/{sender_id}", data=data, follow_redirects=True)
        assert saved.status_code == 200
        with portal.app_context():
            sender = db.session.get(InvitationSender, sender_id)
            for field in ["male_template", "female_template", "plural_template"]:
                assert getattr(sender, field) == data[field]
            assert sender.access_token == token
            assert db.session.get(InvitationSender, 3).male_template == "{name} היקר"
            assert db.session.scalar(db.select(db.func.count(InvitationSender.id))) == 3
    assert prepare(admin).json["text"] == "משה זכר 1"
    assert prepare(admin, guest_id=2, prefix="/u/bride-token").json["text"] == "מרים נקבה 2"


@pytest.mark.parametrize(
    "query,included,excluded",
    [
        ("q=משה", "משה", "מרים"),
        ("side=bride", "מרים", "משה"),
        ("group=חברים", "משה", "מרים"),
    ],
)
def test_filters(admin, query, included, excluded):
    data = admin.get(ROOT + "/admin?" + query).data.decode()
    assert f"<h3>{included}" in data
    assert f"<h3>{excluded}" not in data


def workbook_file(rows):
    book = Workbook()
    sheet = book.active
    sheet.append(["מזהה", "שם פרטי", "שם משפחה", "טלפון", "צד", "צורת פנייה", "סטטוס"])
    for values in rows:
        sheet.append(values)
    data = BytesIO()
    book.save(data)
    data.seek(0)
    return data


def test_excel_roundtrip_preserves_ids_history_and_phone(admin, portal):
    sent = prepare(admin).json
    finish(admin, sent["attempt"])
    before = row(portal)
    exported = admin.get(ROOT + "/admin/export.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data))
    sheet = workbook["מוזמנים"]
    assert sheet["D2"].value == "0501234567"
    sheet["B2"] = "משה מעודכן"
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    response = admin.post(ROOT + "/admin/import", data={"file": (data, "guests.xlsx")})
    assert response.status_code == 302
    assert row(portal) == before
    with portal.app_context():
        assert db.session.get(InvitationPortalGuest, 1).first_name == "משה מעודכן"
        assert db.session.scalar(db.select(db.func.count(InvitationPortalGuest.id))) == 3


@pytest.mark.parametrize("side,gender", [("חתן", "זכר"), ("bride", "female"), ("groom", "plural")])
def test_import_new_requires_explicit_side_and_gender(admin, portal, side, gender):
    data = workbook_file([["", "new", "", "0501234567", side, gender, ""]])
    assert (
        admin.post(ROOT + "/admin/import", data={"file": (data, "guests.xlsx")}).status_code == 302
    )
    with portal.app_context():
        guest = db.session.scalar(
            db.select(InvitationPortalGuest).where(InvitationPortalGuest.first_name == "new")
        )
        assert guest.side in {"groom", "bride"}
        assert guest.salutation in {"male", "female", "plural"}


@pytest.mark.parametrize(
    "bad_row",
    [
        ["missing-id", "new", "", "", "חתן", "זכר", ""],
        ["", "new", "", "", "חתן", "", ""],
        ["", "new", "", "", "unknown", "זכר", ""],
        ["", "new", "", "", "חתן", "זכר", "wrong"],
    ],
)
def test_bad_import_is_atomic(admin, portal, bad_row):
    data = workbook_file([["", "valid", "", "", "חתן", "זכר", ""], bad_row])
    response = admin.post(ROOT + "/admin/import", data={"file": (data, "guests.xlsx")})
    assert response.status_code == 400
    with portal.app_context():
        assert db.session.scalar(db.select(db.func.count(InvitationPortalGuest.id))) == 3


def test_sender_excel_is_scoped_and_rejects_other_side(client, portal):
    exported = client.get(ROOT + "/u/groom-token/export.xlsx")
    sheet = load_workbook(BytesIO(exported.data))["מוזמנים"]
    assert sheet.max_row == 3
    data = workbook_file([["", "new", "", "", "כלה", "נקבה", ""]])
    response = client.post(ROOT + "/u/groom-token/import", data={"file": (data, "guests.xlsx")})
    assert response.status_code == 400


@pytest.mark.parametrize("name,raw", [("bad.xlsx", b"not-a-zip"), ("bad.exe", b"x")])
def test_corrupt_upload_returns_helpful_error(admin, name, raw):
    assert (
        admin.post(ROOT + "/admin/import", data={"file": (BytesIO(raw), name)}).status_code == 400
    )


def test_image_upload_and_invalid_image(admin):
    response = admin.post(
        ROOT + "/admin/settings", data={"image": (BytesIO(b"not an image"), "test.jpg")}
    )
    assert response.status_code == 200
    assert "אינו תמונה".encode() in response.data
    image = BytesIO()
    Image.new("RGB", (20, 30), "white").save(image, "PNG")
    image.seek(0)
    assert (
        admin.post(ROOT + "/admin/settings", data={"image": (image, "test.png")}).status_code == 302
    )
    response = admin.get(ROOT + "/image")
    assert response.status_code == 200
    assert response.mimetype == "image/jpeg"


def test_excel_formula_injection_exported_as_literal(admin, portal):
    with portal.app_context():
        db.session.get(InvitationPortalGuest, 1).first_name = '=HYPERLINK("bad")'
        db.session.commit()
    sheet = load_workbook(BytesIO(admin.get(ROOT + "/admin/export.xlsx").data))["מוזמנים"]
    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value.startswith("=")


def test_csrf_enforced_and_legacy_login_does_not_authorize_portal(client, portal):
    portal.config["WTF_CSRF_ENABLED"] = True
    assert client.post(ROOT + "/u/groom-token/guest/1/prepare").status_code == 400
    with client.session_transaction() as session:
        session["_user_id"] = "1"
    assert client.get(ROOT + "/admin").status_code == 302


def test_inactive_admin_denied(admin, portal):
    with portal.app_context():
        db.session.get(User, 1).is_active_account = False
        db.session.commit()
    assert admin.get(ROOT + "/admin").status_code == 302
    assert prepare(admin, prefix="/admin").status_code == 401


def test_init_db_is_repeatable_and_preserves_portal_guests(portal):
    runner = portal.test_cli_runner()
    assert runner.invoke(args=["init-db"]).exit_code == 0
    assert runner.invoke(args=["init-db"]).exit_code == 0
    with portal.app_context():
        assert db.session.scalar(db.select(db.func.count(InvitationPortalGuest.id))) == 3
        assert db.session.scalar(db.select(db.func.count(InvitationSendAttempt.id))) == 0


def test_sender_edit_invalidates_prepared_text(client, portal):
    data = prepare(client).json
    response = client.post(
        ROOT + "/u/groom-token/guest/1/edit",
        data={
            "first_name": "משה מעודכן",
            "side": "groom",
            "salutation": "plural",
            "phone": "0501234567",
        },
    )
    assert response.status_code == 302
    assert finish(client, data["attempt"]).status_code == 409
    assert prepare(client).json["text"] == "משה מעודכן היקרים"


def test_resend_counts_only_after_confirmation(client, portal):
    first = prepare(client).json
    finish(client, first["attempt"])
    second = prepare(client).json
    assert row(portal)[1] == 1
    finish(client, second["attempt"])
    assert row(portal)[:2] == ("sent", 2)


def test_status_filter(admin):
    data = prepare(admin).json
    finish(admin, data["attempt"])
    response = admin.get(ROOT + "/admin?status=sent")
    assert "<h3>משה לוי".encode() in response.data
    assert "<h3>מרים כהן".encode() not in response.data


def test_duplicate_excel_ids_rejected(admin, portal):
    with portal.app_context():
        identifier = db.session.get(InvitationPortalGuest, 1).external_id
    values = [identifier, "משה", "לוי", "0501234567", "חתן", "זכר", ""]
    data = workbook_file([values, values])
    assert admin.post(ROOT + "/admin/import", data={"file": (data, "g.xlsx")}).status_code == 400


def test_csrf_valid_token_accepts_send(client, portal):
    import re

    portal.config["WTF_CSRF_ENABLED"] = True
    html = client.get(ROOT + "/u/groom-token").data.decode()
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)
    assert prepare(client, csrf_token=csrf).status_code == 200


def test_logout_revokes_portal_access(admin):
    assert admin.post(ROOT + "/admin/logout").status_code == 302
    assert admin.get(ROOT + "/admin").status_code == 302


def test_template_substitution_is_literal_and_escaped_in_html(client, portal):
    with portal.app_context():
        sender = db.session.get(InvitationSender, 1)
        sender.male_template = "{name}: {unexpected} and {name.__class__}"
        db.session.get(InvitationPortalGuest, 1).first_name = "<script>alert(1)</script>"
        db.session.commit()
    page = client.get(ROOT + "/u/groom-token").data.decode()
    assert "<script>alert(1)</script>" not in page
    data = prepare(client).json
    assert "{unexpected}" in data["text"] and "{name.__class__}" in data["text"]


def test_privacy_headers(client, portal):
    response = client.get(ROOT + "/u/groom-token")
    assert response.headers["Cache-Control"] == "no-store"
    assert response.headers["Referrer-Policy"] == "same-origin"
    assert b'<meta name="referrer" content="same-origin">' in response.data


@pytest.mark.parametrize("operation", ["prepare", "edit"])
def test_https_csrf_requires_same_origin_referrer(client, portal, operation):
    import re

    portal.config.update(WTF_CSRF_ENABLED=True, WTF_CSRF_SSL_STRICT=True)
    origin = "https://portal.example.test"
    page = ROOT + "/u/groom-token"
    html = client.get(page, base_url=origin).data.decode()
    csrf = re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)
    path = page + "/guest/1/" + operation
    data = {"csrf_token": csrf}
    if operation == "edit":
        data.update(first_name="משה", side="groom", salutation="male", phone="0501234567")
    missing = client.post(path, base_url=origin, data=data)
    assert missing.status_code == 400
    assert b"referrer header is missing" in missing.data
    foreign = client.post(
        path, base_url=origin, data=data, headers={"Referer": "https://other.test/"}
    )
    assert foreign.status_code == 400
    good = client.post(path, base_url=origin, data=data, headers={"Referer": origin + page})
    assert good.status_code == (200 if operation == "prepare" else 302)

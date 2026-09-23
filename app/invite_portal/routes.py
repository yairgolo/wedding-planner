from __future__ import annotations

import csv
import secrets
from datetime import datetime, timezone
from io import BytesIO, StringIO
from pathlib import Path

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, UnidentifiedImageError
from sqlalchemy import or_
from werkzeug.security import check_password_hash

from app.extensions import db, limiter
from app.models import (
    InvitationPortalActivity,
    InvitationPortalGuest,
    InvitationPortalSettings,
    InvitationSender,
    User,
)

from .forms import ImageForm, PortalGuestForm, PortalLoginForm, SenderForm

invite_portal_bp = Blueprint("invite_portal", __name__, url_prefix="/invite-manager")

DEFAULT_TEMPLATES = {
    "male": (
        "{name} היקר,\n\nנשמח מאוד להזמינך לחגוג איתנו ביום המיוחד שלנו."
        "\nנשמח לראותך איתנו. 🤍"
    ),
    "female": (
        "{name} היקרה,\n\nנשמח מאוד להזמינך לחגוג איתנו ביום המיוחד שלנו."
        "\nנשמח לראותך איתנו. 🤍"
    ),
    "plural": (
        "{name} היקרים,\n\nנשמח מאוד להזמינכם לחגוג איתנו ביום המיוחד שלנו."
        "\nנשמח לראותכם איתנו. 🤍"
    ),
}
HEADERS = [
    "מזהה",
    "שם פרטי",
    "שם משפחה",
    "טלפון",
    "צד",
    "צורת פנייה",
    "קבוצה",
    "סטטוס",
    "נשלח על ידי",
    "נשלח בתאריך",
    "ניסיונות",
    "הערה",
]
SIDE_LABELS = {"groom": "חתן", "bride": "כלה"}
SALUTATION_LABELS = {"male": "זכר", "female": "נקבה", "plural": "רבים"}
STATUS_LABELS = {"unsent": "טרם נשלח", "preparing": "בטיפול", "sent": "נשלח"}
SIDE_VALUES = {**SIDE_LABELS, "חתן": "groom", "כלה": "bride"}
SALUTATION_VALUES = {**SALUTATION_LABELS, "זכר": "male", "נקבה": "female", "רבים": "plural"}
STATUS_VALUES = {**STATUS_LABELS, "טרם נשלח": "unsent", "בטיפול": "preparing", "נשלח": "sent"}


def now() -> datetime:
    return datetime.now(timezone.utc)


def clean_phone(value: str | None) -> str | None:
    cleaned = "".join(ch for ch in (value or "") if ch.isdigit() or ch == "+")
    return cleaned or None


def settings() -> InvitationPortalSettings:
    item = db.session.scalar(db.select(InvitationPortalSettings).limit(1))
    if not item:
        item = InvitationPortalSettings()
        db.session.add(item)
        db.session.commit()
    return item


def admin_id() -> int | None:
    return session.get("invite_portal_admin_id")


def require_admin() -> None:
    user_id = admin_id()
    user = db.session.get(User, user_id) if user_id else None
    if not user or not user.is_admin or not user.is_active:
        abort(403)


def sender_for_token(token: str) -> InvitationSender:
    sender = db.session.scalar(
        db.select(InvitationSender).where(InvitationSender.access_token == token)
    )
    if not sender or not sender.is_active:
        abort(404)
    return sender


def readable_message(sender: InvitationSender, guest: InvitationPortalGuest) -> str:
    template = getattr(sender, f"{guest.salutation}_template")
    try:
        return template.format(name=guest.first_name).strip()
    except (KeyError, ValueError):
        return template.replace("{name}", guest.first_name).strip()


def visible_guests(sender: InvitationSender | None = None):
    stmt = db.select(InvitationPortalGuest)
    if sender:
        stmt = stmt.where(InvitationPortalGuest.side == sender.side)
    q = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()
    if q:
        like = f"%{q}%"
        stmt = stmt.where(
            or_(
                InvitationPortalGuest.first_name.ilike(like),
                InvitationPortalGuest.last_name.ilike(like),
                InvitationPortalGuest.phone.ilike(like),
            )
        )
    if status in STATUS_LABELS:
        stmt = stmt.where(InvitationPortalGuest.status == status)
    return db.session.scalars(
        stmt.order_by(InvitationPortalGuest.status, InvitationPortalGuest.first_name)
    ).all()


def guest_or_404(guest_id: int, sender: InvitationSender | None = None) -> InvitationPortalGuest:
    guest = db.get_or_404(InvitationPortalGuest, guest_id)
    if sender and guest.side != sender.side:
        abort(404)
    return guest


def activity(guest: InvitationPortalGuest, sender: InvitationSender | None, action: str) -> None:
    db.session.add(
        InvitationPortalActivity(
            guest_id=guest.id, sender_id=sender.id if sender else None, action=action
        )
    )


@invite_portal_bp.route("/admin/login", methods=["GET", "POST"])
@limiter.limit("10 per minute")
def admin_login():
    if admin_id():
        return redirect(url_for("invite_portal.admin_index"))
    form = PortalLoginForm()
    if form.validate_on_submit():
        user = db.session.scalar(
            db.select(User).where(User.email == form.email.data.strip().lower())
        )
        if (
            user
            and user.is_admin
            and user.is_active
            and check_password_hash(user.password_hash, form.password.data)
        ):
            session.clear()
            session["invite_portal_admin_id"] = user.id
            session.permanent = True
            return redirect(url_for("invite_portal.admin_index"))
        form.password.errors.append("פרטי הכניסה אינם נכונים.")
    return render_template("invite_portal/login.html", form=form)


@invite_portal_bp.post("/admin/logout")
def admin_logout():
    session.pop("invite_portal_admin_id", None)
    return redirect(url_for("invite_portal.admin_login"))


@invite_portal_bp.get("/")
def home():
    return redirect(
        url_for("invite_portal.admin_index") if admin_id() else url_for("invite_portal.admin_login")
    )


@invite_portal_bp.get("/admin")
def admin_index():
    require_admin()
    guests = visible_guests()
    all_guests = db.session.scalars(db.select(InvitationPortalGuest)).all()
    stats = {key: sum(g.status == key for g in all_guests) for key in STATUS_LABELS}
    return render_template(
        "invite_portal/dashboard.html",
        admin=True,
        sender=None,
        guests=guests,
        stats=stats,
        settings=settings(),
        labels=(SIDE_LABELS, SALUTATION_LABELS, STATUS_LABELS),
    )


@invite_portal_bp.route("/admin/guest/new", methods=["GET", "POST"])
@invite_portal_bp.route("/admin/guest/<int:guest_id>/edit", methods=["GET", "POST"])
def admin_guest_form(guest_id: int | None = None):
    require_admin()
    guest = (
        db.get_or_404(InvitationPortalGuest, guest_id)
        if guest_id
        else InvitationPortalGuest(side="groom", salutation="male")
    )
    form = PortalGuestForm(obj=guest)
    if form.validate_on_submit():
        form.populate_obj(guest)
        guest.phone = clean_phone(form.phone.data)
        db.session.add(guest)
        db.session.commit()
        flash("המוזמן נשמר.", "success")
        return redirect(url_for("invite_portal.admin_index"))
    return render_template(
        "invite_portal/guest_form.html", form=form, admin=True, sender=None, guest=guest
    )


@invite_portal_bp.post("/admin/guest/<int:guest_id>/status")
def admin_guest_status(guest_id: int):
    require_admin()
    guest = db.get_or_404(InvitationPortalGuest, guest_id)
    status = request.form.get("status")
    if status not in STATUS_LABELS:
        abort(400)
    guest.status = status
    if status != "sent":
        guest.sent_at = None
    activity(guest, None, f"status_{status}")
    db.session.commit()
    return redirect(url_for("invite_portal.admin_index", **request.args))


@invite_portal_bp.route("/admin/senders", methods=["GET", "POST"])
@invite_portal_bp.route("/admin/senders/<int:sender_id>", methods=["GET", "POST"])
def admin_senders(sender_id: int | None = None):
    require_admin()
    sender = (
        db.get_or_404(InvitationSender, sender_id)
        if sender_id
        else InvitationSender(side="groom", **DEFAULT_TEMPLATES)
    )
    form = SenderForm(obj=sender)
    if form.validate_on_submit():
        form.populate_obj(sender)
        db.session.add(sender)
        db.session.commit()
        flash("פרופיל השולח נשמר. הקישור האישי מופיע לידו.", "success")
        return redirect(url_for("invite_portal.admin_senders"))
    senders = db.session.scalars(
        db.select(InvitationSender).order_by(InvitationSender.side, InvitationSender.name)
    ).all()
    return render_template(
        "invite_portal/senders.html",
        form=form,
        editing=sender_id is not None,
        sender=sender,
        senders=senders,
        admin=True,
        side_labels=SIDE_LABELS,
    )


@invite_portal_bp.post("/admin/senders/<int:sender_id>/renew")
def renew_sender_link(sender_id: int):
    require_admin()
    sender = db.get_or_404(InvitationSender, sender_id)
    sender.access_token = secrets.token_urlsafe(32)
    db.session.commit()
    flash("נוצר קישור חדש. הקישור הקודם בוטל.", "success")
    return redirect(url_for("invite_portal.admin_senders"))


@invite_portal_bp.route("/admin/settings", methods=["GET", "POST"])
def admin_settings():
    require_admin()
    item = settings()
    form = ImageForm()
    if form.validate_on_submit() and form.image.data:
        try:
            image = Image.open(form.image.data.stream)
            image.verify()
            form.image.data.stream.seek(0)
            image = Image.open(form.image.data.stream).convert("RGB")
        except (UnidentifiedImageError, OSError):
            form.image.errors.append("הקובץ אינו תמונה תקינה.")
        else:
            filename = f"invite-portal-{secrets.token_hex(8)}.jpg"
            target = Path(current_app.config["UPLOAD_FOLDER"]) / filename
            image.thumbnail((2400, 3200))
            image.save(target, "JPEG", quality=92, optimize=True)
            if item.image_filename:
                (Path(current_app.config["UPLOAD_FOLDER"]) / item.image_filename).unlink(
                    missing_ok=True
                )
            item.image_filename = filename
            db.session.commit()
            flash("תמונת ההזמנה עודכנה.", "success")
            return redirect(url_for("invite_portal.admin_settings"))
    return render_template("invite_portal/settings.html", admin=True, form=form, settings=item)


@invite_portal_bp.get("/image")
def image():
    item = settings()
    if not item.image_filename:
        abort(404)
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], item.image_filename)


@invite_portal_bp.get("/admin/export.xlsx")
def export_excel():
    require_admin()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "מוזמנים"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="805D66")
    for guest in db.session.scalars(
        db.select(InvitationPortalGuest).order_by(InvitationPortalGuest.id)
    ):
        sheet.append(
            [
                guest.external_id,
                guest.first_name,
                guest.last_name or "",
                guest.phone or "",
                SIDE_LABELS[guest.side],
                SALUTATION_LABELS[guest.salutation],
                guest.group_name or "",
                STATUS_LABELS[guest.status],
                guest.last_sender.name if guest.last_sender else "",
                guest.sent_at.strftime("%d/%m/%Y %H:%M") if guest.sent_at else "",
                guest.attempts,
                guest.notes or "",
            ]
        )
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(
            max(len(str(c.value or "")) for c in column) + 2, 32
        )
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    return send_file(
        data,
        as_attachment=True,
        download_name="invitation-manager-guests.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def rows_from_upload(upload):
    if (upload.filename or "").lower().endswith(".xlsx"):
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook["מוזמנים"] if "מוזמנים" in workbook.sheetnames else workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, ())
        return [dict(zip(headers, row, strict=False)) for row in values]
    raw = upload.stream.read().decode("utf-8-sig", errors="replace")
    return list(csv.DictReader(StringIO(raw)))


@invite_portal_bp.post("/admin/import")
def import_excel():
    require_admin()
    upload = request.files.get("file")
    if not upload or not upload.filename or not upload.filename.lower().endswith((".xlsx", ".csv")):
        flash("יש לבחור קובץ Excel או CSV.", "error")
        return redirect(url_for("invite_portal.admin_index"))
    created = updated = skipped = 0
    for row in rows_from_upload(upload):
        external_id = str(row.get("מזהה") or "").strip()
        guest = (
            db.session.scalar(
                db.select(InvitationPortalGuest).where(
                    InvitationPortalGuest.external_id == external_id
                )
            )
            if external_id
            else None
        )
        first_name = str(row.get("שם פרטי") or "").strip()
        if not first_name:
            skipped += 1
            continue
        guest = guest or InvitationPortalGuest(side="groom", salutation="male")
        guest.first_name = first_name[:120]
        guest.last_name = str(row.get("שם משפחה") or "").strip()[:120] or None
        guest.phone = clean_phone(str(row.get("טלפון") or ""))
        guest.side = SIDE_VALUES.get(str(row.get("צד") or "").strip(), "groom")
        guest.salutation = SALUTATION_VALUES.get(str(row.get("צורת פנייה") or "").strip(), "male")
        guest.group_name = str(row.get("קבוצה") or "").strip()[:120] or None
        guest.status = STATUS_VALUES.get(str(row.get("סטטוס") or "").strip(), "unsent")
        guest.notes = str(row.get("הערה") or "").strip() or None
        db.session.add(guest)
        if guest.id:
            updated += 1
        else:
            created += 1
    db.session.commit()
    flash(f"הייבוא הסתיים: {created} נוספו, {updated} עודכנו, {skipped} דולגו.", "success")
    return redirect(url_for("invite_portal.admin_index"))


@invite_portal_bp.get("/u/<token>")
def user_index(token: str):
    sender = sender_for_token(token)
    guests = visible_guests(sender)
    all_guests = db.session.scalars(
        db.select(InvitationPortalGuest).where(InvitationPortalGuest.side == sender.side)
    ).all()
    stats = {key: sum(g.status == key for g in all_guests) for key in STATUS_LABELS}
    return render_template(
        "invite_portal/dashboard.html",
        admin=False,
        sender=sender,
        guests=guests,
        stats=stats,
        settings=settings(),
        labels=(SIDE_LABELS, SALUTATION_LABELS, STATUS_LABELS),
    )


@invite_portal_bp.route("/u/<token>/guest/new", methods=["GET", "POST"])
@invite_portal_bp.route("/u/<token>/guest/<int:guest_id>/edit", methods=["GET", "POST"])
def user_guest_form(token: str, guest_id: int | None = None):
    sender = sender_for_token(token)
    guest = (
        guest_or_404(guest_id, sender)
        if guest_id
        else InvitationPortalGuest(side=sender.side, salutation="male")
    )
    form = PortalGuestForm(obj=guest)
    form.side.data = sender.side
    if form.validate_on_submit():
        form.populate_obj(guest)
        guest.side = sender.side
        guest.phone = clean_phone(form.phone.data)
        db.session.add(guest)
        db.session.commit()
        flash("המוזמן נשמר.", "success")
        return redirect(url_for("invite_portal.user_index", token=token))
    return render_template(
        "invite_portal/guest_form.html", form=form, admin=False, sender=sender, guest=guest
    )


@invite_portal_bp.post("/u/<token>/guest/<int:guest_id>/prepare")
@limiter.limit("120 per hour")
def prepare_send(token: str, guest_id: int):
    sender = sender_for_token(token)
    guest = guest_or_404(guest_id, sender)
    if not guest.phone:
        return jsonify({"ok": False, "error": "חסר מספר טלפון למוזמן."}), 400
    guest.status = "preparing"
    activity(guest, sender, "preparing")
    db.session.commit()
    return jsonify(
        {
            "ok": True,
            "text": readable_message(sender, guest),
            "image_url": url_for("invite_portal.image") if settings().image_filename else None,
        }
    )


@invite_portal_bp.post("/u/<token>/guest/<int:guest_id>/confirm")
@limiter.limit("120 per hour")
def confirm_send(token: str, guest_id: int):
    sender = sender_for_token(token)
    guest = guest_or_404(guest_id, sender)
    if request.form.get("sent") == "true":
        guest.status = "sent"
        guest.sent_at = now()
        guest.last_sender_id = sender.id
        guest.attempts += 1
        activity(guest, sender, "sent")
    else:
        guest.status = "unsent"
        activity(guest, sender, "cancelled")
    db.session.commit()
    return jsonify({"ok": True})


@invite_portal_bp.post("/u/<token>/guest/<int:guest_id>/status")
def user_guest_status(token: str, guest_id: int):
    sender = sender_for_token(token)
    guest = guest_or_404(guest_id, sender)
    status = request.form.get("status")
    if status not in STATUS_LABELS:
        abort(400)
    guest.status = status
    if status != "sent":
        guest.sent_at = None
    activity(guest, sender, f"status_{status}")
    db.session.commit()
    return redirect(url_for("invite_portal.user_index", token=token))

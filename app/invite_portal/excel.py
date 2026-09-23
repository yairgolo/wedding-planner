from __future__ import annotations

import csv
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from flask import flash, redirect, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from app.extensions import db
from app.models import InvitationPortalGuest

from .routes import (
    DECISION_LABELS,
    SALUTATION_LABELS,
    SIDE_LABELS,
    STATUS_LABELS,
    clean_phone,
    close_pending,
    index_url,
    invalidate_attempts,
    invite_portal_bp,
    now,
    require_admin,
    sender_for_token,
)

HEADERS = [
    "מזהה",
    "שם פרטי",
    "שם משפחה",
    "טלפון",
    "צד",
    "צורת פנייה",
    "קבוצה",
    "סטטוס",
    "נשלח על ידי",
    "נשלח בתאריך",
    "ניסיונות",
    "הערה",
    "החלטת הזמנה",
]


def enum_value(value, labels):
    raw = str(value or "").strip()
    return raw if raw in labels else {v: k for k, v in labels.items()}.get(raw)


def export(sender=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "מוזמנים"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "C2"
    sheet.append(HEADERS)
    query = (
        db.select(InvitationPortalGuest)
        .where(InvitationPortalGuest.deleted_at.is_(None))
        .order_by(InvitationPortalGuest.id)
    )
    if sender:
        query = query.where(InvitationPortalGuest.side == sender.side)
    for guest in db.session.scalars(query):
        sheet.append(
            [
                guest.external_id,
                guest.first_name,
                guest.last_name or "",
                guest.phone or "",
                SIDE_LABELS.get(guest.side, ""),
                SALUTATION_LABELS.get(guest.salutation, ""),
                guest.group_name or "",
                STATUS_LABELS[guest.status],
                guest.last_sender.name if guest.last_sender else "",
                guest.sent_at.strftime("%d/%m/%Y %H:%M") if guest.sent_at else "",
                guest.attempts,
                guest.notes or "",
                DECISION_LABELS[guest.invitation_decision],
            ]
        )
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # User text must never become an Excel formula.
            cell.alignment = Alignment(horizontal="right", vertical="center")
        sheet.row_dimensions[row[0].row].height = 26
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334B60")
    for letter, width in zip(
        "ABCDEFGHIJKLM", [38, 25, 22, 20, 14, 16, 22, 18, 22, 24, 14, 40, 22], strict=True
    ):
        sheet.column_dimensions[letter].width = width
    for col, labels in (("E", SIDE_LABELS), ("F", SALUTATION_LABELS), ("H", STATUS_LABELS)):
        validation = DataValidation(
            type="list", formula1='"' + ",".join(labels.values()) + '"', allow_blank=True
        )
        validation.errorTitle = "יש לבחור ערך מהרשימה"
        validation.error = "בחרו אחת מהאפשרויות המוצעות."
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(f"{col}2:{col}5001")
    sheet.auto_filter.ref = sheet.dimensions
    guide = workbook.create_sheet("הנחיות")
    for index, label in enumerate(DECISION_LABELS.values(), 1):
        guide.cell(index, 3, label)
    decision_validation = DataValidation(
        type="list", formula1="'הנחיות'!$C$1:$C$2", allow_blank=True
    )
    decision_validation.showErrorMessage = True
    sheet.add_data_validation(decision_validation)
    decision_validation.add("M2:M5001")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 110
    for text in [
        "רשומות קיימות מעודכנות לפי מזהה בלבד. אין לשנות מזהים.",
        "להוספת מוזמן חדש: הוסיפו שורה והשאירו את המזהה ריק.",
        "אפשר לייבא רשומות חלקיות. לפני שליחה חובה להשלים שם פרטי, צד וצורת פנייה.",
        "בעדכון לפי מזהה, שדות ריקים או עמודות חסרות שומרים את הערכים הקיימים.",
        "בייבוא אישי צד חסר משויך לצד של השולח; בייבוא מנהל נשאר ללא שיוך.",
        "טלפון יש להזין כטקסט, כולל האפס בתחילת המספר.",
        "תאריך שליחה, השולח ומספר הניסיונות הם מידע בלבד ונשמרים במערכת.",
        "סטטוס ריק משאיר את הסטטוס הקיים. שינוי סטטוס סוגר ניסיון שליחה פתוח.",
        "שורה חסרה בקובץ אינה מוחקת מוזמן. הייבוא מוגבל ל־5,000 שורות.",
        "כאשר יש שגיאות לא נשמר אף שינוי; מתקנים ומעלים שוב.",
        "החלטת הזמנה: כן, מזמינים או בסימן שאלה. בסימן שאלה השליחה חסומה.",
        "החלטה ריקה משאירה את ההחלטה הקיימת. מוזמן חדש מוגדר כן, מזמינים.",
        "מוזמן שנמחק אינו מיוצא; לשחזור משתמשים בעמוד מוזמנים שנמחקו באתר.",
    ]:
        guide.append([text])
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    return send_file(
        data,
        as_attachment=True,
        download_name="invitation-manager-guests.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def read_rows(upload):
    if upload.filename.lower().endswith(".xlsx"):
        with ZipFile(upload.stream) as archive:
            if sum(f.file_size for f in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("קובץ Excel גדול מדי.")
        upload.stream.seek(0)
        workbook = load_workbook(upload.stream, read_only=True, data_only=False)
        try:
            sheet = workbook["מוזמנים"] if "מוזמנים" in workbook.sheetnames else workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(x or "").strip() for x in next(values, ())]
            rows = []
            for index, row in enumerate(values):
                if index >= 5000:
                    raise ValueError("אפשר לייבא עד 5,000 שורות בכל פעם.")
                rows.append(dict(zip(headers, row, strict=False)))
            return headers, rows
        finally:
            workbook.close()
    text = upload.stream.read().decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    rows = []
    for index, row in enumerate(reader):
        if index >= 5000:
            raise ValueError("אפשר לייבא עד 5,000 שורות בכל פעם.")
        rows.append(row)
    return reader.fieldnames or [], rows


def import_file(sender=None):
    upload = request.files.get("file")
    errors, changes, seen = [], [], set()
    if not upload or not upload.filename.lower().endswith((".xlsx", ".csv")):
        errors.append("יש לבחור קובץ Excel (.xlsx) או CSV.")
    else:
        try:
            headers, rows = read_rows(upload)
            if not set(HEADERS).intersection(headers):
                raise ValueError("לא נמצאו עמודות מוכרות.")
            for number, row in enumerate(rows, 2):
                if not any(value is not None and str(value).strip() for value in row.values()):
                    continue
                identifier = str(row.get("מזהה") or "").strip()
                guest = (
                    db.session.scalar(
                        db.select(InvitationPortalGuest).where(
                            InvitationPortalGuest.external_id == identifier
                        )
                    )
                    if identifier
                    else None
                )

                def field(header, attribute, default="", row=row, guest=guest):
                    raw = str(row.get(header) or "").strip()
                    return raw or (getattr(guest, attribute) or "" if guest else default)

                first = field("שם פרטי", "first_name")
                side_raw = str(row.get("צד") or "").strip()
                salutation_raw = str(row.get("צורת פנייה") or "").strip()
                side = (
                    enum_value(side_raw, SIDE_LABELS)
                    if side_raw
                    else field("צד", "side", sender.side if sender else "")
                )
                salutation = (
                    enum_value(salutation_raw, SALUTATION_LABELS)
                    if salutation_raw
                    else field("צורת פנייה", "salutation")
                )
                status_raw = str(row.get("סטטוס") or "").strip()
                status = (
                    enum_value(status_raw, STATUS_LABELS)
                    if status_raw
                    else (guest.status if guest else "unsent")
                )
                decision_raw = str(row.get("החלטת הזמנה") or "").strip()
                decision = (
                    enum_value(decision_raw, DECISION_LABELS)
                    if decision_raw
                    else (guest.invitation_decision if guest else "invited")
                )
                error = None
                if identifier and (not guest or identifier in seen):
                    error = "מזהה לא מוכר או כפול בקובץ. לשורה חדשה יש להשאיר מזהה ריק."
                elif sender and (side != sender.side or (guest and guest.side != sender.side)):
                    error = "אפשר לייבא ולעדכן רק מוזמנים מהצד שלך."
                elif guest and guest.deleted_at:
                    error = "המוזמן נמחק. יש לשחזר אותו באתר לפני עדכון דרך Excel."
                elif len(first) > 120:
                    error = "השם הפרטי ארוך מדי."
                elif side is None or salutation is None or not status or not decision:
                    error = "יש לבחור צד, צורת פנייה, סטטוס והחלטת הזמנה תקינים."
                elif (
                    (decision == "undecided" or not salutation or not first or not side)
                    and status in {"preparing", "sent"}
                    and (not guest or status != guest.status)
                ):
                    error = "מוזמן בסימן שאלה או עם פרטים חסרים אינו יכול לעבור לבטיפול או לנשלח."
                if error:
                    errors.append(f"שורה {number}: {error}")
                    continue
                seen.add(identifier)
                changes.append(
                    (
                        guest,
                        dict(
                            first_name=first,
                            last_name=field("שם משפחה", "last_name")[:120],
                            phone=clean_phone(field("טלפון", "phone")),
                            side=side,
                            salutation=salutation,
                            group_name=field("קבוצה", "group_name")[:120],
                            status=status,
                            invitation_decision=decision,
                            notes=field("הערה", "notes"),
                        ),
                    )
                )
        except (
            ValueError,
            BadZipFile,
            OSError,
            KeyError,
            UnicodeError,
            SyntaxError,
            InvalidFileException,
        ):
            errors.append(
                "לא ניתן לקרוא את הקובץ. ודאו שהוא Excel תקין עם עמודות החובה ועד 5,000 שורות."
            )
    if errors:
        return render_template(
            "invite_portal/import_result.html", admin=sender is None, sender=sender, errors=errors
        ), 400
    created = updated = 0
    for guest, values in changes:
        if guest:
            updated += 1
            if values["invitation_decision"] == "undecided" or not all(
                values[key] for key in ("first_name", "side", "salutation")
            ):
                close_pending(guest)
                if values["status"] == "preparing":
                    values["status"] = guest.status
            else:
                invalidate_attempts(guest.id)
            if values["status"] != guest.status:
                guest.sent_at = now() if values["status"] == "sent" else None
                guest.last_sender_id = None
        else:
            created += 1
            guest = InvitationPortalGuest(sent_at=now() if values["status"] == "sent" else None)
        for key, value in values.items():
            setattr(guest, key, value)
        db.session.add(guest)
    db.session.commit()
    flash(f"הייבוא הסתיים: {created} נוספו, {updated} עודכנו.", "success")
    return redirect(index_url(sender))


@invite_portal_bp.get("/admin/export.xlsx")
def export_excel():
    require_admin()
    return export()


@invite_portal_bp.post("/admin/import")
def import_excel():
    require_admin()
    return import_file()


@invite_portal_bp.get("/u/<token>/export.xlsx")
def user_export_excel(token):
    return export(sender_for_token(token))


@invite_portal_bp.post("/u/<token>/import")
def user_import_excel(token):
    return import_file(sender_for_token(token))
